import streamlit as st
import pandas as pd
import re
import io
import requests
import concurrent.futures
import time
import threading
from typing import Optional, Dict, Any
import unicodedata
import xml.etree.ElementTree as ET
from datetime import datetime
from rapidfuzz import fuzz, process

# ==============================================================================
# CONFIGURAÇÃO DA PÁGINA E ESTILOS
# ==============================================================================
st.set_page_config(
    page_title="Portal de Compliance & Controles Internos",
    page_icon="🛡️",
    layout="wide"
)

# ==============================================================================
# CONSTANTES
# ==============================================================================
API_PEP_URL = "https://api.portaldatransparencia.gov.br/api-de-dados/peps"
TIMEOUT_API_SEGUNDOS = 20
MAX_THREADS_API = 8  # ↓ reduzido para evitar bloqueio/rajadas

# ==============================================================================
# FUNÇÕES UTILITÁRIAS (HELPER FUNCTIONS)
# ==============================================================================

def get_api_key() -> str:
    """Lê API KEY com fallback e faz strip() para remover \n e espaços invisíveis."""
    try:
        k = st.secrets["api_keys"]["API_KEY"]
    except Exception:
        k = st.secrets.get("API_KEY", "")
    return (k or "").strip()

API_KEY = get_api_key()

DEFAULT_HEADERS = {
    "chave-api-dados": API_KEY,
    "Accept": "application/json",
    "User-Agent": "Mozilla/5.0 (Windows NT 10.0; Win64; x64) StreamlitCompliance/1.0",
}

_thread_local = threading.local()

def get_session() -> requests.Session:
    """Session por thread para reaproveitar conexões (menos chance de bloqueio)."""
    if not hasattr(_thread_local, "session"):
        _thread_local.session = requests.Session()
    return _thread_local.session

def sanitizar_documento(valor: Any) -> Optional[str]:
    """Remove tudo que não é dígito. Retorna None se vazio/nulo."""
    if pd.isna(valor):
        return None
    limpo = re.sub(r"\D", "", str(valor))
    return limpo if limpo else None

def normalizar_cpf_padrao(valor: Any) -> Optional[str]:
    """
    Normalização robusta de CPF:
    - Trata float vindo do Excel (remove .0)
    - Remove não numéricos
    - Preenche com zeros À ESQUERDA até 11 dígitos
    """
    if pd.isna(valor):
        return None

    # ⚠️ CASO CRÍTICO: CPF veio como número (float) do Excel
    if isinstance(valor, float):
        # remove a parte decimal SEM criar zeros artificiais
        valor = str(int(valor))

    # agora sim trata como string
    limpo = re.sub(r"\D", "", str(valor))

    if not limpo:
        return None

    # completa à ESQUERDA, nunca à direita
    return limpo.zfill(11)


def extrair_numero_contrato(texto: str) -> Optional[str]:
    if isinstance(texto, str):
        match_regex = re.search(r'Contrato:\s*(\d+)\s*\(', texto)
        return match_regex.group(1) if match_regex else None
    return None

def extrair_nome_empreendimento(texto: str) -> Optional[str]:
    if isinstance(texto, str):
        match_regex = re.search(r'^\d{11} - ', texto)
        return texto.split(' - ')[1] if match_regex else None
    return None

def extrair_id_fornecedor(texto: str) -> Optional[str]:
    if isinstance(texto, str):
        match_regex = re.search(r'Fornecedor:\s*(\d+)\s*-', texto)
        return match_regex.group(1) if match_regex else None
    return None

def consultar_api_pep(cpf: str) -> str:
    """
    Consulta CPF na API PEP com robustez:
    - CPF normalizado (zfill 11)
    - retry/backoff para 403/405/429/5xx
    - sem redirects
    - força Connection: close para reduzir instabilidade em CloudFront com rajadas
    """
    cpf_tratado = normalizar_cpf_padrao(cpf)

    if not cpf_tratado:
        return "CPF Não Informado"
    if cpf_tratado == "00000000000":
        return "CPF Zerado (Inválido)"
    if len(cpf_tratado) != 11:
        return "CPF Inválido (Tam. Incorreto)"

    if not API_KEY:
        return "Erro: API_KEY não configurada"

    # headers por requisição (evita mutação/efeitos colaterais)
    headers = {
        **DEFAULT_HEADERS,
        "Connection": "close",  # ← importante p/ reduzir bug/instabilidade com CF
    }

    tentativas = 4
    for i in range(tentativas):
        try:
            # NÃO reutilizar session aqui é mais estável em alguns ambientes
            r = requests.get(
                API_PEP_URL,
                headers=headers,
                params={"cpf": cpf_tratado, "pagina": 1},
                timeout=TIMEOUT_API_SEGUNDOS,
                allow_redirects=False
            )

            # Se teve redirect, isso é pista forte (e evita "seguir" p/ lugar errado)
            if r.status_code in (301, 302, 307, 308):
                loc = r.headers.get("Location", "")
                return f"Erro: Redirect {r.status_code} -> {loc}"

            if r.status_code == 200:
                return "Sim - PEP Identificado" if r.json() else "Não consta"

            if r.status_code == 404:
                return "Não consta"

            # 👇 trate 405 como instabilidade também
            if r.status_code in (403, 405, 429, 500, 502, 503, 504):
                time.sleep((0.8 * (2 ** i)) + 0.05)  # backoff
                continue

            if r.status_code == 401:
                return "Erro: API_KEY inválida (401)"

            # Para qualquer outro, guarda um pedacinho do body (útil p/ diagnóstico)
            body_hint = (r.text or "")[:120].replace("\n", " ").strip()
            return f"Erro HTTP {r.status_code} - {body_hint}"

        except requests.RequestException:
            time.sleep((0.8 * (2 ** i)) + 0.05)
            continue

    return "Erro: instabilidade/limite (CloudFront)"


def processar_parser_relatorio_complexo(arquivo_obj) -> pd.DataFrame:
    lista_dados_estruturados = []

    try:
        df_bruto = pd.read_excel(arquivo_obj, header=None)
    except Exception as e:
        st.error(f"Falha na leitura do Excel: {e}")
        return pd.DataFrame()

    empresa_atual = None
    contrato_atual = None
    data_contrato_atual = None

    for _, linha in df_bruto.iterrows():
        primeira_celula = str(linha[0]).strip() if pd.notna(linha[0]) else ""

        if "Empreendimento:" in primeira_celula:
            if len(linha) > 1 and pd.notna(linha[1]):
                empresa_atual = str(linha[1]).strip()
            continue

        if primeira_celula.isdigit():
            if len(linha) >= 5:
                contrato_atual = str(linha[3]).strip() if pd.notna(linha[3]) else ""
                data_bruta = linha[4]
                if isinstance(data_bruta, pd.Timestamp):
                    data_contrato_atual = data_bruta.strftime('%Y-%m-%d')
                else:
                    data_contrato_atual = str(data_bruta).strip() if pd.notna(data_bruta) else ""

                nome_cliente = str(linha[1]).strip() if pd.notna(linha[1]) else ""

                lista_dados_estruturados.append({
                    "Numero Contrato": contrato_atual,
                    "Empreendimento": empresa_atual,
                    "Nome da Parte": nome_cliente,
                    "Tipo Parte": "CLIENTE PRINCIPAL",
                    "CPF": None,
                    "Data Cadastro": data_contrato_atual
                })
            continue

        if "Cônjuge:" in primeira_celula:
            nome_conjuge = str(linha[1]).strip() if (len(linha) > 1 and pd.notna(linha[1])) else ""
            lista_dados_estruturados.append({
                "Numero Contrato": contrato_atual,
                "Empreendimento": empresa_atual,
                "Nome da Parte": nome_conjuge,
                "Tipo Parte": "CONJUGE",
                "CPF": None,
                "Data Cadastro": data_contrato_atual
            })
            continue

        if "Participante:" in primeira_celula:
            nome_participante = str(linha[1]).strip() if (len(linha) > 1 and pd.notna(linha[1])) else ""
            lista_dados_estruturados.append({
                "Numero Contrato": contrato_atual,
                "Empreendimento": empresa_atual,
                "Nome da Parte": nome_participante,
                "Tipo Parte": "PARTICIPANTE",
                "CPF": None,
                "Data Cadastro": data_contrato_atual
            })
            continue

        if primeira_celula.startswith("CPF"):
            if lista_dados_estruturados:
                cpf_valor = str(linha[1]).strip() if (len(linha) > 1 and pd.notna(linha[1])) else ""
                lista_dados_estruturados[-1]["CPF"] = cpf_valor
            continue

    df_estruturado = pd.DataFrame(lista_dados_estruturados)
    if df_estruturado.empty:
        return df_estruturado

    df_estruturado = df_estruturado[df_estruturado['Nome da Parte'] != ""]
    termos_ignorar = ["Clientes no Bloco", "Clientes no Empreendimento", "Clientes", "Total Clientes", "Total Geral"]
    df_estruturado = df_estruturado[~df_estruturado['Nome da Parte'].astype(str).str.strip().isin(termos_ignorar)]

    df_estruturado['cpf_normalizado'] = df_estruturado['CPF'].apply(normalizar_cpf_padrao)

    mask_cpf_problematico = (
        (df_estruturado['cpf_normalizado'].isnull()) |
        (df_estruturado['cpf_normalizado'] == '00000000000')
    )

    df_invalidos = df_estruturado[mask_cpf_problematico]
    df_validos = df_estruturado[~mask_cpf_problematico]

    df_validos = df_validos.drop_duplicates(subset=['cpf_normalizado'], keep='first')

    df_final = pd.concat([df_validos, df_invalidos], ignore_index=True)
    return df_final

# ==============================================================================
# LISTAS RESTRITIVAS (OFAC / ONU / INTERPOL / FBI)
# ==============================================================================
LR_LIMIAR_PADRAO = 88            # similaridade mínima (0-100) para possível match
LR_TIMEOUT = 90
LR_PAUSA_INTERPOL = 1.0          # segundos entre consultas à Interpol (evita bloqueio)
LR_CACHE_TTL = 24 * 3600         # listas baixadas ficam em cache por 24h

LR_HEADERS = {
    "User-Agent": ("Mozilla/5.0 (Windows NT 10.0; Win64; x64) AppleWebKit/537.36 "
                   "(KHTML, like Gecko) Chrome/124.0 Safari/537.36"),
    "Accept": "application/json, text/xml, text/csv, */*",
}

LR_OFAC_BASE = "https://sanctionslistservice.ofac.treas.gov/api/PublicationPreview/exports/"
LR_URLS = {
    "ofac_sdn": [LR_OFAC_BASE + "SDN.CSV",
                 "https://www.treasury.gov/ofac/downloads/sdn.csv"],
    "ofac_sdn_alt": [LR_OFAC_BASE + "ALT.CSV",
                     "https://www.treasury.gov/ofac/downloads/alt.csv"],
    "ofac_cons": [LR_OFAC_BASE + "CONS_PRIM.CSV",
                  "https://www.treasury.gov/ofac/downloads/consolidated/cons_prim.csv"],
    "ofac_cons_alt": [LR_OFAC_BASE + "CONS_ALT.CSV",
                      "https://www.treasury.gov/ofac/downloads/consolidated/cons_alt.csv"],
    "onu": ["https://scsanctions.un.org/resources/xml/en/consolidated.xml"],
}
LR_INTERPOL_URL = "https://ws-public.interpol.int/notices/v1/red"
LR_FBI_URL = "https://api.fbi.gov/wanted/v1/list"

LR_STOPWORDS = {"DA", "DE", "DO", "DAS", "DOS", "E"}
LR_FONTES = {"ofac": "OFAC", "onu": "ONU", "interpol": "Interpol", "fbi": "FBI"}

LR_COLS_OFAC_PRIM = ["ent_num", "nome", "tipo", "programa", "titulo", "call_sign", "vess_type",
                     "tonnage", "grt", "vess_flag", "vess_owner", "remarks"]
LR_COLS_OFAC_ALT = ["ent_num", "alt_num", "alt_type", "alt_name", "alt_remarks"]


def lr_agora() -> str:
    return datetime.now().strftime("%d/%m/%Y %H:%M:%S")


def lr_normalizar(texto: Any) -> str:
    """Maiúsculas, sem acento, sem pontuação, sem preposições (DA/DE/DO...)."""
    if texto is None or (isinstance(texto, float) and pd.isna(texto)):
        return ""
    t = unicodedata.normalize("NFKD", str(texto))
    t = "".join(c for c in t if not unicodedata.combining(c))
    t = re.sub(r"[^A-Za-z0-9]+", " ", t).upper()
    return " ".join(p for p in t.split() if p not in LR_STOPWORDS)


class LrIndiceLista:
    """Índice de nomes de uma lista para busca aproximada (fuzzy)."""

    def __init__(self, registros):
        self.registros = [r for r in registros if lr_normalizar(r.get("nome_na_lista"))]
        self.chaves = [lr_normalizar(r["nome_na_lista"]) for r in self.registros]

    def buscar(self, nome, limiar):
        q = lr_normalizar(nome)
        if not q:
            return []
        achados = process.extract(q, self.chaves, scorer=fuzz.token_sort_ratio,
                                  score_cutoff=limiar, limit=None)
        melhores = {}
        for _, score, idx in achados:
            reg = dict(self.registros[idx])
            reg["score"] = round(score, 1)
            chave = (reg["fonte"], reg["lista"], reg["id_na_lista"])
            if chave not in melhores or reg["score"] > melhores[chave]["score"]:
                melhores[chave] = reg
        return sorted(melhores.values(), key=lambda r: -r["score"])


def _lr_baixar(chave: str):
    """Baixa o arquivo tentando as URLs alternativas. Levanta erro se todas falharem."""
    erros = []
    for url in LR_URLS[chave]:
        try:
            r = requests.get(url, headers=LR_HEADERS, timeout=LR_TIMEOUT)
            r.raise_for_status()
            if len(r.content) < 1000:
                raise ValueError("arquivo vazio ou inválido")
            return r.content, url
        except Exception as e:  # noqa: BLE001
            erros.append(f"{url}: {e}")
    raise RuntimeError(" | ".join(erros))


def _lr_ler_csv_ofac(conteudo: bytes, colunas) -> pd.DataFrame:
    df = pd.read_csv(io.BytesIO(conteudo), header=None, names=colunas,
                     usecols=range(len(colunas)), dtype=str, encoding="latin-1",
                     on_bad_lines="skip", engine="python")
    df = df.apply(lambda s: s.str.strip())
    return df.replace({"-0-": ""}).fillna("")


@st.cache_data(ttl=LR_CACHE_TTL, show_spinner=False)
def lr_carregar_ofac():
    """SDN + Consolidated (Non-SDN), com aliases. Qualquer falha invalida a fonte toda."""
    registros, log = [], []
    for prim, alt, rotulo in (("ofac_sdn", "ofac_sdn_alt", "SDN"),
                              ("ofac_cons", "ofac_cons_alt", "Consolidated (Non-SDN)")):
        conteudo, url = _lr_baixar(prim)
        dp = _lr_ler_csv_ofac(conteudo, LR_COLS_OFAC_PRIM)
        dp = dp[dp["ent_num"].str.fullmatch(r"\d+")]
        info = dp.set_index("ent_num")[["nome", "tipo", "programa", "remarks"]].to_dict("index")

        def novo(ent, nome, tipo_nome, d):
            return {
                "fonte": "OFAC", "lista": rotulo, "id_na_lista": ent,
                "nome_na_lista": nome, "nome_principal": d["nome"],
                "tipo_nome": tipo_nome, "tipo_registro": d["tipo"] or "entity",
                "programa": d["programa"], "detalhe": d["remarks"][:300],
                "link": f"https://sanctionssearch.ofac.treas.gov/Details.aspx?id={ent}",
            }

        for ent, d in info.items():
            registros.append(novo(ent, d["nome"], "principal", d))

        conteudo_alt, _ = _lr_baixar(alt)
        da = _lr_ler_csv_ofac(conteudo_alt, LR_COLS_OFAC_ALT)
        for ent, nome_alt, tipo_alt in zip(da["ent_num"], da["alt_name"], da["alt_type"]):
            d = info.get(ent)
            if d and nome_alt:
                registros.append(novo(ent, nome_alt, f"alias ({tipo_alt})", d))

        log.append({"fonte": f"OFAC {rotulo}", "url": url, "data_hora_download": lr_agora(),
                    "status": "OK", "qtd_registros": len(dp), "qtd_aliases": len(da),
                    "observacao": ""})
    return registros, log


@st.cache_data(ttl=LR_CACHE_TTL, show_spinner=False)
def lr_carregar_onu():
    conteudo, url = _lr_baixar("onu")
    root = ET.fromstring(conteudo)
    registros = []
    qtd = 0

    def add(nome, ref, lista, tipo_nome, tipo_reg, principal, comentario):
        if nome and nome.strip():
            registros.append({
                "fonte": "ONU", "lista": f"UN Consolidated - {lista}",
                "id_na_lista": ref, "nome_na_lista": nome.strip(),
                "nome_principal": principal, "tipo_nome": tipo_nome,
                "tipo_registro": tipo_reg, "programa": lista,
                "detalhe": (comentario or "")[:300],
                "link": "https://main.un.org/securitycouncil/en/content/un-sc-consolidated-list",
            })

    for ind in root.iter("INDIVIDUAL"):
        qtd += 1
        partes = [ind.findtext(t) for t in ("FIRST_NAME", "SECOND_NAME", "THIRD_NAME", "FOURTH_NAME")]
        nome = " ".join(p.strip() for p in partes if p and p.strip())
        ref = ind.findtext("REFERENCE_NUMBER") or ind.findtext("DATAID") or ""
        lista = ind.findtext("UN_LIST_TYPE") or ""
        com = ind.findtext("COMMENTS1") or ""
        add(nome, ref, lista, "principal", "Individual", nome, com)
        add(ind.findtext("NAME_ORIGINAL_SCRIPT"), ref, lista, "escrita original", "Individual", nome, com)
        for al in ind.findall("INDIVIDUAL_ALIAS"):
            add(al.findtext("ALIAS_NAME"), ref, lista,
                f"alias ({al.findtext('QUALITY') or ''})", "Individual", nome, com)

    for ent in root.iter("ENTITY"):
        qtd += 1
        nome = (ent.findtext("FIRST_NAME") or "").strip()
        ref = ent.findtext("REFERENCE_NUMBER") or ent.findtext("DATAID") or ""
        lista = ent.findtext("UN_LIST_TYPE") or ""
        com = ent.findtext("COMMENTS1") or ""
        add(nome, ref, lista, "principal", "Entity", nome, com)
        for al in ent.findall("ENTITY_ALIAS"):
            add(al.findtext("ALIAS_NAME"), ref, lista,
                f"alias ({al.findtext('QUALITY') or ''})", "Entity", nome, com)

    if qtd == 0:
        raise RuntimeError("XML da ONU sem registros (formato pode ter mudado)")
    log = [{"fonte": "ONU - Consolidated List", "url": url, "data_hora_download": lr_agora(),
            "status": "OK", "qtd_registros": qtd, "qtd_aliases": len(registros) - qtd,
            "observacao": ""}]
    return registros, log


@st.cache_data(ttl=LR_CACHE_TTL, show_spinner=False)
def lr_carregar_fbi():
    itens, pagina = [], 1
    while True:
        r = requests.get(LR_FBI_URL, params={"page": pagina, "pageSize": 50},
                         headers=LR_HEADERS, timeout=LR_TIMEOUT)
        r.raise_for_status()
        dados = r.json()
        lote = dados.get("items") or []
        if not lote:
            break
        itens.extend(lote)
        if len(itens) >= (dados.get("total") or 0):
            break
        pagina += 1
        time.sleep(0.5)
    if not itens:
        raise RuntimeError("API do FBI não retornou registros")

    registros = []
    for it in itens:
        titulo = it.get("title") or ""
        base = {
            "fonte": "FBI", "lista": "FBI Wanted", "id_na_lista": it.get("uid") or "",
            "nome_principal": titulo, "tipo_registro": "Individual",
            "programa": ", ".join(it.get("subjects") or []),
            "detalhe": re.sub(r"<[^>]+>", " ", it.get("caution") or "")[:300].strip(),
            "link": it.get("url") or "",
        }
        registros.append({**base, "nome_na_lista": titulo, "tipo_nome": "principal"})
        for alias in it.get("aliases") or []:
            registros.append({**base, "nome_na_lista": alias, "tipo_nome": "alias"})

    log = [{"fonte": "FBI Wanted", "url": LR_FBI_URL, "data_hora_download": lr_agora(),
            "status": "OK", "qtd_registros": len(itens),
            "qtd_aliases": len(registros) - len(itens), "observacao": ""}]
    return registros, log


class LrConsultaInterpol:
    """Consulta online das Red Notices, nome a nome (forename + sobrenome)."""

    def __init__(self):
        self.sessao = requests.Session()
        self.sessao.headers.update({**LR_HEADERS,
                                    "Origin": "https://www.interpol.int",
                                    "Referer": "https://www.interpol.int/"})

    def buscar(self, nome, limiar):
        tokens = lr_normalizar(nome).split()
        if not tokens:
            return []
        params = {"name": tokens[-1], "resultPerPage": 160}
        if len(tokens) > 1:
            params["forename"] = tokens[0]

        ultimo_erro = None
        for tentativa in range(3):
            try:
                r = self.sessao.get(LR_INTERPOL_URL, params=params, timeout=LR_TIMEOUT)
                if r.status_code in (403, 429, 500, 502, 503, 504):
                    raise requests.HTTPError(f"HTTP {r.status_code}")
                r.raise_for_status()
                notices = (r.json().get("_embedded") or {}).get("notices") or []
                break
            except Exception as e:  # noqa: BLE001
                ultimo_erro = e
                time.sleep(LR_PAUSA_INTERPOL * (2 ** tentativa))
        else:
            raise RuntimeError(f"Interpol indisponível: {ultimo_erro}")
        time.sleep(LR_PAUSA_INTERPOL)

        q = lr_normalizar(nome)
        resultados = []
        for n in notices:
            nome_lista = f"{n.get('forename') or ''} {n.get('name') or ''}".strip()
            score = fuzz.token_sort_ratio(q, lr_normalizar(nome_lista))
            if score < limiar:
                continue
            resultados.append({
                "fonte": "INTERPOL", "lista": "Red Notice",
                "id_na_lista": n.get("entity_id") or "",
                "nome_na_lista": nome_lista, "nome_principal": nome_lista,
                "tipo_nome": "principal", "tipo_registro": "Individual",
                "programa": "Nacionalidades: " + ", ".join(n.get("nationalities") or []),
                "detalhe": f"Nascimento: {n.get('date_of_birth') or 'n/d'}",
                "link": ((n.get("_links") or {}).get("self") or {}).get("href", ""),
                "score": round(score, 1),
            })
        return sorted(resultados, key=lambda r: -r["score"])


def lr_ler_planilha(arquivo) -> pd.DataFrame:
    if arquivo.name.lower().endswith(".csv"):
        bruto = arquivo.getvalue()
        for enc in ("utf-8-sig", "latin-1"):
            try:
                return pd.read_csv(io.BytesIO(bruto), sep=None, engine="python", dtype=str, encoding=enc)
            except UnicodeDecodeError:
                continue
    return pd.read_excel(arquivo, dtype=str)


def lr_executar_triagem(nomes, fontes, limiar, barra, status):
    """Retorna (resultado_por_nome, ocorrencias, log_fontes, erro_fonte)."""
    carregadores = {"ofac": lr_carregar_ofac, "onu": lr_carregar_onu, "fbi": lr_carregar_fbi}
    indices, erro_fonte, log_fontes = {}, {}, []

    for f in fontes:
        if f in carregadores:
            status.info(f"🔹 Carregando lista {LR_FONTES[f]}...")
            try:
                registros, log = carregadores[f]()
                indices[f] = LrIndiceLista(registros)
                log_fontes.extend(log)
            except Exception as e:  # noqa: BLE001
                erro_fonte[f] = str(e)
                log_fontes.append({"fonte": LR_FONTES[f], "url": "", "data_hora_download": lr_agora(),
                                   "status": "ERRO", "qtd_registros": None, "qtd_aliases": None,
                                   "observacao": str(e)[:500]})

    interpol = None
    if "interpol" in fontes:
        interpol = LrConsultaInterpol()
        log_fontes.append({"fonte": "INTERPOL Red Notices", "url": LR_INTERPOL_URL,
                           "data_hora_download": lr_agora(), "status": "consulta online",
                           "qtd_registros": None, "qtd_aliases": None,
                           "observacao": "Consulta nome a nome (primeiro nome + sobrenome)"})

    resultado, ocorrencias, erros_interpol = {}, [], 0
    total = len(nomes)
    for i, nome in enumerate(nomes, 1):
        barra.progress(i / total, text=f"Consultando {i}/{total}: {nome}")
        linha = {}
        tem_match = tem_erro = False
        for f in fontes:
            col = f"Status {LR_FONTES[f]}"
            try:
                if f == "interpol":
                    achados = interpol.buscar(nome, limiar)
                elif f in indices:
                    achados = indices[f].buscar(nome, limiar)
                else:
                    raise RuntimeError(erro_fonte.get(f, "fonte indisponível"))
            except Exception as e:  # noqa: BLE001
                linha[col] = "ERRO - consultar manualmente"
                tem_erro = True
                if f == "interpol":
                    erros_interpol += 1
                    ocorrencias.append({"nome_consultado": nome, "fonte": "INTERPOL",
                                        "detalhe": f"ERRO NA CONSULTA: {e}"})
                continue

            if achados:
                tem_match = True
                linha[col] = f"POSSÍVEL MATCH ({len(achados)})"
                linha[f"Maior Score {LR_FONTES[f]}"] = achados[0]["score"]
                ocorrencias.extend({"nome_consultado": nome, **a} for a in achados)
            else:
                linha[col] = "SEM OCORRÊNCIA"

        linha["Resultado Triagem"] = ("ANALISAR" if tem_match
                                      else "INCOMPLETO" if tem_erro else "SEM OCORRÊNCIA")
        resultado[nome] = linha

    if interpol is not None and erros_interpol:
        log_fontes[-1]["observacao"] += f" | {erros_interpol} consulta(s) com erro"
    return resultado, ocorrencias, log_fontes, erro_fonte


def lr_gerar_excel(df_resumo, df_ocorr, df_fontes) -> bytes:
    from openpyxl.styles import Font, PatternFill
    cores = {"ANALISAR": "F8CBAD", "INCOMPLETO": "FFE699", "SEM OCORRÊNCIA": "C6EFCE"}
    buffer = io.BytesIO()
    with pd.ExcelWriter(buffer, engine="openpyxl") as writer:
        for aba, df in (("Resumo", df_resumo), ("Ocorrencias", df_ocorr), ("Fontes", df_fontes)):
            df.to_excel(writer, sheet_name=aba, index=False)
            ws = writer.sheets[aba]
            ws.freeze_panes = "A2"
            ws.auto_filter.ref = ws.dimensions
            for c in ws[1]:
                c.font = Font(bold=True, color="FFFFFF")
                c.fill = PatternFill("solid", fgColor="1F4E78")
            for col in ws.columns:
                largura = max((len(str(c.value)) for c in col if c.value is not None), default=10)
                ws.column_dimensions[col[0].column_letter].width = min(max(largura + 2, 12), 60)
            if aba == "Resumo" and "Resultado Triagem" in df.columns:
                idx = list(df.columns).index("Resultado Triagem") + 1
                for row in ws.iter_rows(min_row=2, min_col=idx, max_col=idx):
                    cor = cores.get(row[0].value)
                    if cor:
                        row[0].fill = PatternFill("solid", fgColor=cor)
    buffer.seek(0)
    return buffer.getvalue()


def lr_gerar_modelo() -> bytes:
    modelo = pd.DataFrame({
        "nome": ["Fulano de Tal", "Empresa Exemplo Ltda"],
        "documento": ["000.000.000-00", "00.000.000/0001-00"],
        "tipo": ["PF", "PJ"],
        "observacao": ["exemplo - apague esta linha", "exemplo - apague esta linha"],
    })
    buffer = io.BytesIO()
    with pd.ExcelWriter(buffer, engine="openpyxl") as writer:
        modelo.to_excel(writer, sheet_name="nomes", index=False)
        ws = writer.sheets["nomes"]
        for col, w in zip("ABCD", (45, 22, 8, 35)):
            ws.column_dimensions[col].width = w
    buffer.seek(0)
    return buffer.getvalue()


# ==============================================================================
# INTERFACE PRINCIPAL
# ==============================================================================
abas_navegacao = st.tabs(["📊 Fornecedores sem Diligência", "🔎 Verificação de PEPs",
                          "🌐 Listas Restritivas"])

# ------------------------------------------------------------------------------
# ABA 1
# ------------------------------------------------------------------------------
with abas_navegacao[0]:
    st.title("📊 Controle Interno - Fornecedores sem Diligência")
    st.markdown("Cruza dados de contratos com a base de *Due Diligence* para identificar contratos com fornecedores não homologados.")
    col1, col2, col3, col4 = st.columns(4)

    with col1:
        arquivo_contratos = st.file_uploader("Relatório de Contratos (.xlsx)", type="xlsx")
    with col2:
        arquivo_itens = st.file_uploader("Relatório de Itens (.xlsx)", type="xlsx")
    with col3:
        arquivo_agentes = st.file_uploader("Base de Agentes (.xlsx)", type="xlsx")
    with col4:
        arquivo_diligencias = st.file_uploader("Controle de Diligências (.xlsx)", type="xlsx")

    if st.button("Executar Cruzamento de Dados", key="btn_cruzamento"):
        if arquivo_contratos and arquivo_itens and arquivo_agentes and arquivo_diligencias:
            try:
                with st.spinner("Normalizando dados e cruzando bases..."):
                    df_contratos = pd.read_excel(arquivo_contratos)
                    df_itens = pd.read_excel(arquivo_itens)
                    df_agentes = pd.read_excel(arquivo_agentes)
                    df_diligencias = pd.read_excel(arquivo_diligencias, sheet_name="CONTROLE DDR")

                    mapa_colunas_erp = {
                        '(a)': 'item_codigo',
                        '(b)': 'descricao_item',
                        '(c)': 'unid_medida',
                        '(d)': 'quantidade',
                        '(e)': 'valor_unitario',
                        '(f)': 'valor_total',
                        '(k = g - j)': 'quantidade_saldo',
                        '(l = i + k)': 'valor_saldo'
                    }
                    df_contratos.rename(columns=mapa_colunas_erp, inplace=True)

                    filtro_lixo = 'Total Geral:|Total do Projeto:|Exibir Apropriações|Parâmetros Selecionados|OCMEG'
                    df_contratos = df_contratos[~df_contratos['item_codigo'].astype(str).str.contains(filtro_lixo, na=False)]

                    df_contratos['Contrato_Ref'] = df_contratos['item_codigo'].apply(extrair_numero_contrato).ffill()
                    df_contratos['Empreendimento_Ref'] = df_contratos['item_codigo'].apply(extrair_nome_empreendimento).ffill()
                    df_contratos['ID_Fornecedor'] = df_contratos['descricao_item'].apply(extrair_id_fornecedor).ffill()

                    cols_descarte = ['(g = e - f)', '(h)', '(i)', '(j)', 'Unnamed: 12', 'Unnamed: 13', 'Unnamed: 14', 'Unnamed: 15', 'Unnamed: 16']
                    df_contratos = df_contratos.drop(columns=[c for c in cols_descarte if c in df_contratos.columns])

                    df_contratos = df_contratos[~df_contratos['item_codigo'].astype(str).str.contains('-', na=False)]

                    df_contratos['ID_Fornecedor'] = df_contratos['ID_Fornecedor'].astype(str)
                    df_agentes['Código'] = df_agentes['Código'].astype(str)

                    df_agentes['CNPJ_Sanitizado'] = df_agentes['CNPJ'].apply(sanitizar_documento)
                    df_diligencias['CNPJ_Sanitizado'] = df_diligencias['CNPJ'].apply(sanitizar_documento)

                    df_diligencias_unicas = df_diligencias[['CNPJ_Sanitizado']].drop_duplicates()

                    df_contratos = pd.merge(
                        df_contratos,
                        df_itens[['Cód.Item', 'Definição Item']],
                        left_on='item_codigo',
                        right_on='Cód.Item',
                        how='left'
                    )

                    df_contratos = pd.merge(
                        df_contratos,
                        df_agentes[['Código', 'CNPJ', 'CNPJ_Sanitizado', 'Nome fantasia']],
                        left_on='ID_Fornecedor',
                        right_on='Código',
                        how='left'
                    )

                    cols_monetarias = ['valor_unitario', 'valor_total', 'valor_saldo', 'quantidade', 'quantidade_saldo']
                    for col in cols_monetarias:
                        df_contratos[col] = pd.to_numeric(df_contratos[col], errors='coerce').fillna(0)

                    cols_group_contrato = ['Contrato_Ref', 'Empreendimento_Ref', 'ID_Fornecedor', 'Nome fantasia', 'CNPJ']
                    df_visao_contratos = df_contratos.groupby(cols_group_contrato)[cols_monetarias].sum().reset_index()
                    df_visao_contratos.drop(columns=['valor_unitario', 'quantidade_saldo', 'quantidade'], inplace=True)

                    cols_group_fornecedor = ['ID_Fornecedor', 'Nome fantasia', 'CNPJ', 'CNPJ_Sanitizado']
                    df_visao_fornecedores = df_contratos.groupby(cols_group_fornecedor)[cols_monetarias].sum().reset_index()

                    filtro_servicos = df_contratos[df_contratos['Definição Item'] == 'SE']
                    soma_apenas_servicos = filtro_servicos.groupby('ID_Fornecedor')['valor_total'].sum().reset_index()
                    soma_apenas_servicos.rename(columns={'valor_total': 'Total_Apenas_Servicos'}, inplace=True)

                    df_visao_fornecedores = pd.merge(
                        df_visao_fornecedores,
                        soma_apenas_servicos,
                        on='ID_Fornecedor',
                        how='left'
                    )
                    df_visao_fornecedores['Total_Apenas_Servicos'] = df_visao_fornecedores['Total_Apenas_Servicos'].fillna(0)

                    df_visao_fornecedores['Flag_Risco_20k'] = df_visao_fornecedores['Total_Apenas_Servicos'].apply(
                        lambda x: '⚠️ Atinge 20 mil' if x >= 20000 else 'Baixo Valor'
                    )

                    df_visao_fornecedores = pd.merge(
                        df_visao_fornecedores,
                        df_diligencias_unicas,
                        on='CNPJ_Sanitizado',
                        how='left',
                        indicator=True
                    )

                    df_visao_fornecedores['Status_Diligencia'] = df_visao_fornecedores['_merge'].apply(
                        lambda x: '✅ Realizada' if x == 'both' else '❌ Pendente'
                    )

                    df_visao_fornecedores.drop(
                        columns=['_merge', 'CNPJ_Sanitizado', 'valor_unitario', 'quantidade_saldo', 'quantidade'],
                        inplace=True
                    )
                    df_contratos.drop(columns=['CNPJ_Sanitizado'], inplace=True, errors='ignore')

                    buffer_excel = io.BytesIO()
                    with pd.ExcelWriter(buffer_excel, engine='openpyxl') as writer:
                        df_contratos.to_excel(writer, sheet_name='Analitico_Detalhado', index=False)
                        df_visao_contratos.to_excel(writer, sheet_name='Sintetico_Contratos', index=False)
                        df_visao_fornecedores.to_excel(writer, sheet_name='Matriz_Compliance', index=False)

                    buffer_excel.seek(0)

                st.success("Base processada com sucesso!")
                st.download_button(
                    label="📥 Baixar Relatório de Compliance Consolidado",
                    data=buffer_excel,
                    file_name="Relatorio_Compliance_Fornecedores.xlsx",
                    mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
                )

                st.markdown("### 📋 Matriz de Riscos (Fornecedores)")
                st.dataframe(df_visao_fornecedores.head(10))

            except Exception as e:
                st.error(f"Erro no processamento dos arquivos: {e}")
        else:
            st.warning("⚠️ Atenção: Todos os 4 arquivos são obrigatórios para o cruzamento de dados.")

# ------------------------------------------------------------------------------
# ABA 2: VERIFICAÇÃO DE PEPs
# ------------------------------------------------------------------------------
with abas_navegacao[1]:
    st.title("🔎 Consulta Unificada de PEPs")

    # --------------------------------------------------------------------------
    # SEÇÃO 1: INCORPORADORA
    # --------------------------------------------------------------------------
    st.subheader("🏢 Incorporadora (Relatório ERP)")
    arquivo_bruto_pep = st.file_uploader("Upload Relatório Incorporadora (.xlsx)", type=["xlsx"], key="upload_incorp")

    if st.button("Verificar Base Incorporadora", key="btn_incorp"):
        if arquivo_bruto_pep is None:
            st.warning("Por favor, faça o upload do arquivo.")
        else:
            try:
                st.info("🔹 Fase 1/2: Estruturando dados e normalizando CPFs (Adicionando zeros)...")
                df_clientes_tratado = processar_parser_relatorio_complexo(arquivo_bruto_pep)

                if df_clientes_tratado.empty:
                    st.error("O arquivo está vazio ou fora do padrão.")
                else:
                    mask_consulta = (
                        df_clientes_tratado['cpf_normalizado'].notna()
                        & (df_clientes_tratado['cpf_normalizado'] != '00000000000')
                    )
                    lista_cpfs_consulta = df_clientes_tratado.loc[mask_consulta, "cpf_normalizado"].tolist()

                    st.info(f"🔹 Fase 2/2: Consultando {len(lista_cpfs_consulta)} CPFs na API...")

                    with concurrent.futures.ThreadPoolExecutor(max_workers=MAX_THREADS_API) as executor:
                        resultados_raw = list(executor.map(consultar_api_pep, lista_cpfs_consulta))

                    mapa_resultados = dict(zip(lista_cpfs_consulta, resultados_raw))

                    df_clientes_tratado["Status_PEP"] = df_clientes_tratado["cpf_normalizado"].map(mapa_resultados)
                    df_clientes_tratado.loc[df_clientes_tratado["cpf_normalizado"].isna(), "Status_PEP"] = "CPF Inválido/Não Identificado"

                    st.success("✅ Verificação Incorporadora Finalizada!")
                    st.dataframe(df_clientes_tratado)

                    buffer_excel_pep = io.BytesIO()
                    with pd.ExcelWriter(buffer_excel_pep, engine='openpyxl') as writer:
                        df_export = df_clientes_tratado.drop(columns=['cpf_normalizado'], errors='ignore')
                        df_export.to_excel(writer, index=False, sheet_name="Analise_PEP")
                    buffer_excel_pep.seek(0)

                    st.download_button(
                        "📥 Baixar Relatório",
                        data=buffer_excel_pep,
                        file_name="Resultado_PEP_Incorporadora.xlsx",
                        mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
                    )

            except Exception as e:
                st.error(f"Erro crítico: {e}")

    st.markdown("---")

    # --------------------------------------------------------------------------
    # SEÇÃO 2: AUTOMOTIVO
    # --------------------------------------------------------------------------
    st.subheader("🚗 Automotivo (Planilha Simples)")
    arquivo_bruto_auto = st.file_uploader("Upload Base Automotiva (.xlsx)", type=["xlsx"], key="upload_auto")

    if st.button("Verificar Base Automotiva", key="btn_auto"):
        if arquivo_bruto_auto is None:
            st.warning("Upload necessário.")
        else:
            try:
                st.info("🔹 Lendo arquivo e normalizando dados (CPF com zero à esquerda até 11 dígitos)...")
                df_auto = pd.read_excel(arquivo_bruto_auto)
                df_auto.columns = [c.upper().strip() for c in df_auto.columns]

                if 'CPF' not in df_auto.columns:
                    st.error("Coluna CPF não encontrada.")
                else:
                    df_auto["CPF_Normalizado"] = df_auto["CPF"].apply(normalizar_cpf_padrao)

                    mask_cpf_valido = df_auto["CPF_Normalizado"].notna() & (df_auto["CPF_Normalizado"] != "00000000000")
                    cpfs_para_consulta = df_auto.loc[mask_cpf_valido, "CPF_Normalizado"].drop_duplicates().tolist()

                    st.info(f"🔹 Consultando {len(cpfs_para_consulta)} CPFs únicos na API...")

                    with concurrent.futures.ThreadPoolExecutor(max_workers=MAX_THREADS_API) as executor:
                        resultados_raw = list(executor.map(consultar_api_pep, cpfs_para_consulta))

                    mapa_resultados = dict(zip(cpfs_para_consulta, resultados_raw))

                    df_auto["Status_PEP"] = df_auto["CPF_Normalizado"].map(mapa_resultados)
                    df_auto.loc[~mask_cpf_valido, "Status_PEP"] = "CPF Inválido/Não Identificado"

                    st.success("✅ Verificação Finalizada!")
                    st.dataframe(df_auto)

                    buffer_excel_auto = io.BytesIO()
                    with pd.ExcelWriter(buffer_excel_auto, engine='openpyxl') as writer:
                        df_export = df_auto.drop(columns=['CPF_Normalizado'], errors='ignore')
                        df_export.to_excel(writer, index=False, sheet_name="Analise_PEP_Auto")
                    buffer_excel_auto.seek(0)

                    st.download_button(
                        "📥 Baixar Relatório",
                        data=buffer_excel_auto,
                        file_name="Resultado_PEP_Automotivo.xlsx",
                        mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
                    )

            except Exception as e:
                st.error(f"Erro: {e}")

# ------------------------------------------------------------------------------
# ABA 3: LISTAS RESTRITIVAS (OFAC / ONU / INTERPOL / FBI)
# ------------------------------------------------------------------------------
with abas_navegacao[2]:
    st.title("🌐 Triagem em Listas Restritivas")
    st.markdown(
        "Consulta uma lista de nomes (pessoas ou empresas) na **OFAC** (SDN e Non-SDN), "
        "na **Lista Consolidada da ONU**, nas **Red Notices da Interpol** e na lista de "
        "procurados do **FBI**. A comparação é aproximada: ignora acentos, preposições e a ordem das palavras."
    )

    st.download_button(
        "📄 Baixar planilha modelo",
        data=lr_gerar_modelo(),
        file_name="modelo_triagem_nomes.xlsx",
        mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        key="lr_modelo",
    )

    arquivo_nomes = st.file_uploader("Planilha com os nomes (.xlsx ou .csv)",
                                     type=["xlsx", "csv"], key="upload_listas")

    col_f, col_l = st.columns(2)
    with col_f:
        fontes_sel = st.multiselect("Fontes", options=list(LR_FONTES), default=list(LR_FONTES),
                                    format_func=LR_FONTES.get, key="lr_fontes")
    with col_l:
        limiar_sel = st.slider("Similaridade mínima (%)", 70, 100, LR_LIMIAR_PADRAO, key="lr_limiar",
                               help="Abaixo de 85 aumenta muito os homônimos; acima de 95 pode perder "
                                    "variações de grafia.")

    df_nomes, coluna_nome = None, None
    if arquivo_nomes is not None:
        try:
            df_nomes = lr_ler_planilha(arquivo_nomes)
        except Exception as e:
            st.error(f"Falha na leitura do arquivo: {e}")
        if df_nomes is not None and not df_nomes.empty:
            colunas = list(df_nomes.columns)
            padrao = next((i for i, c in enumerate(colunas)
                           if str(c).strip().lower() in ("nome", "nomes", "razão social",
                                                         "razao social", "name")), 0)
            coluna_nome = st.selectbox("Coluna com os nomes", colunas, index=padrao, key="lr_coluna")
            qtd_unicos = df_nomes[coluna_nome].dropna().astype(str).str.strip().replace("", pd.NA).dropna().nunique()
            st.caption(f"{len(df_nomes)} linha(s) | {qtd_unicos} nome(s) único(s) para consultar")
            if "interpol" in fontes_sel and qtd_unicos > 0:
                st.caption(f"⏱️ A Interpol é consultada nome a nome: estimativa de "
                           f"~{max(1, round(qtd_unicos * 1.5 / 60))} min.")

    if st.button("Executar Triagem", key="btn_listas"):
        if df_nomes is None or coluna_nome is None:
            st.warning("Faça o upload da planilha com os nomes.")
        elif not fontes_sel:
            st.warning("Selecione ao menos uma fonte.")
        else:
            nomes_serie = df_nomes[coluna_nome].fillna("").astype(str).str.strip()
            nomes_unicos = [n for n in dict.fromkeys(nomes_serie) if n]
            if not nomes_unicos:
                st.warning("A coluna selecionada não tem nomes preenchidos.")
            else:
                status_box = st.empty()
                barra = st.progress(0.0, text="Preparando...")
                try:
                    resultado, ocorrencias, log_fontes, erro_fonte = lr_executar_triagem(
                        nomes_unicos, fontes_sel, limiar_sel, barra, status_box)

                    # Resumo: planilha original + colunas da triagem
                    cols_status = [f"Status {LR_FONTES[f]}" for f in fontes_sel]
                    cols_score = [f"Maior Score {LR_FONTES[f]}" for f in fontes_sel]
                    df_triagem = pd.DataFrame.from_dict(resultado, orient="index")
                    df_triagem = df_triagem.reindex(columns=["Resultado Triagem"] + cols_status + cols_score)
                    df_triagem = df_triagem.dropna(axis=1, how="all")
                    df_resumo = pd.concat(
                        [df_nomes.reset_index(drop=True),
                         df_triagem.reindex(nomes_serie.values).reset_index(drop=True)],
                        axis=1)
                    df_resumo["Resultado Triagem"] = df_resumo["Resultado Triagem"].fillna("NOME VAZIO")

                    cols_ocorr = ["nome_consultado", "fonte", "lista", "score", "nome_na_lista",
                                  "tipo_nome", "nome_principal", "tipo_registro", "programa",
                                  "id_na_lista", "detalhe", "link"]
                    df_ocorr = pd.DataFrame(ocorrencias, columns=cols_ocorr)
                    if not df_ocorr.empty:
                        df_ocorr = df_ocorr.sort_values(["nome_consultado", "score"],
                                                        ascending=[True, False], na_position="last")
                    df_fontes = pd.DataFrame(log_fontes)
                    df_fontes["limiar_similaridade"] = limiar_sel
                    df_fontes["data_hora_triagem"] = lr_agora()

                    st.session_state["lr_resultado"] = {
                        "resumo": df_resumo,
                        "ocorrencias": df_ocorr,
                        "fontes": df_fontes,
                        "erros": erro_fonte,
                        "excel": lr_gerar_excel(df_resumo, df_ocorr, df_fontes),
                        "arquivo": f"Triagem_Listas_Restritivas_{datetime.now():%Y%m%d_%H%M}.xlsx",
                    }
                    status_box.empty()
                    barra.empty()
                except Exception as e:
                    st.error(f"Erro na triagem: {e}")

    res = st.session_state.get("lr_resultado")
    if res:
        st.markdown("---")
        contagem = res["resumo"]["Resultado Triagem"].value_counts()
        m1, m2, m3, m4 = st.columns(4)
        m1.metric("Linhas analisadas", len(res["resumo"]))
        m2.metric("🔴 Analisar", int(contagem.get("ANALISAR", 0)))
        m3.metric("🟡 Incompleto", int(contagem.get("INCOMPLETO", 0)))
        m4.metric("🟢 Sem ocorrência", int(contagem.get("SEM OCORRÊNCIA", 0)))

        for f, erro in res["erros"].items():
            st.warning(f"⚠️ {LR_FONTES[f]} não pôde ser consultada. Os nomes ficaram como "
                       f"'ERRO - consultar manualmente'. Detalhe: {erro[:300]}")
        if int(contagem.get("INCOMPLETO", 0)) and "interpol" not in res["erros"]:
            st.info("Linhas 'INCOMPLETO' tiveram alguma consulta com erro (geralmente a Interpol). "
                    "Veja a aba Ocorrencias do Excel e consulte esses nomes manualmente.")

        st.download_button(
            "📥 Baixar Relatório de Triagem",
            data=res["excel"],
            file_name=res["arquivo"],
            mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
            key="lr_download",
        )

        st.markdown("### 🔎 Possíveis matches")
        matches = res["ocorrencias"][res["ocorrencias"]["score"].notna()]
        if matches.empty:
            st.success("Nenhum possível match encontrado nas fontes consultadas.")
        else:
            st.dataframe(matches, hide_index=True,
                         column_config={"link": st.column_config.LinkColumn("link")})
        st.caption("Match por nome é indício, não confirmação: valide cada caso com data de nascimento, "
                   "nacionalidade ou documento antes de concluir.")

        with st.expander("Resumo por nome"):
            st.dataframe(res["resumo"], hide_index=True)
        with st.expander("Fontes consultadas (trilha de auditoria)"):
            st.dataframe(res["fontes"], hide_index=True)






